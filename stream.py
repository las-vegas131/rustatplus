import streamlit as st
import pandas as pd
import io
import os
import plotly.graph_objects as go
from collections import defaultdict

from config import (
    MIN_MINUTES, SETTINGS_FILE, SELECTED_METRICS_FILE,
    ALL_POSSIBLE_METRICS, DEFAULT_METRICS_WEIGHTS, DEFAULT_MATCH_WEIGHTS,
    METRIC_NAMES_RU, MATCH_METRIC_NAMES_RU, MATCH_ALL_METRICS
)
from utils import (
    load_settings, save_settings, load_selected_metrics, save_selected_metrics,
    load_match_selected_metrics, save_match_selected_metrics
)
from db import (
    check_db_connection, get_leagues, get_seasons_for_leagues,
    get_teams_for_leagues_seasons, get_teams_for_league, get_matches_for_league,
    load_from_db, load_match_stats, get_league_averages, get_player_season_stats_df,
    import_season_excel, import_match_excel, get_db_config
)
from calculations import (
    calculate_ratings, calculate_match_ratings,
    build_position_tables, build_main_table,
    build_match_position_tables, build_match_main_table,
    get_position_group
)
from visualization import (
    create_player_radar_figure, create_compare_figure, create_position_radar,
    export_matches_advanced
)

st.set_page_config(page_title="InStat Analyst", layout="wide")

# CSS для закрепления заголовков и прозрачного фона
st.markdown("""
<style>
thead tr th {
    position: sticky !important;
    top: 0 !important;
    background-color: #f0f2f6 !important;
    z-index: 1;
}
.match-table tbody tr, .match-table tbody td {
    background-color: transparent !important;
}
.stDataFrame tbody tr, .stDataFrame tbody td {
    background-color: transparent !important;
}
</style>
""", unsafe_allow_html=True)

st.title("Анализ футболистов InStat")

# Инициализация session_state
if 'df_db' not in st.session_state:
    st.session_state.df_db = None
if 'df_matches' not in st.session_state:
    st.session_state.df_matches = {}
if 'position_tables' not in st.session_state:
    st.session_state.position_tables = {}
if 'current_settings' not in st.session_state:
    saved = load_settings(SETTINGS_FILE)
    if saved is not None:
        st.session_state.current_settings = saved
    else:
        st.session_state.current_settings = {pos: w.copy() for pos, w in DEFAULT_METRICS_WEIGHTS.items()}
if 'match_settings' not in st.session_state:
    st.session_state.match_settings = {pos: w.copy() for pos, w in DEFAULT_MATCH_WEIGHTS.items()}
if 'selected_main_metrics' not in st.session_state:
    saved = load_selected_metrics(SELECTED_METRICS_FILE)
    st.session_state.selected_main_metrics = saved if saved is not None else []
if 'match_selected_metrics' not in st.session_state:
    saved = load_match_selected_metrics()
    st.session_state.match_selected_metrics = saved if saved is not None else []
if 'avg_source' not in st.session_state:
    st.session_state.avg_source = 'Текущие данные'
if 'avg_league' not in st.session_state:
    st.session_state.avg_league = None
if 'avg_season' not in st.session_state:
    st.session_state.avg_season = None

# Боковая панель (импорт, загрузка и т.д.) – без изменений, как в предыдущей версии
# ... (весь код sidebar остаётся как в последней рабочей версии)

# Основные вкладки
tab_season, tab_match = st.tabs(["📈 Сезон", "⚽ Матч"])

with tab_season:
    if st.session_state.df_db is not None:
        df_active = st.session_state.df_db
        position_tables_active = st.session_state.position_tables
        all_metrics = [m for m in ALL_POSSIBLE_METRICS if m in df_active.columns]
        # Обязательные ТТД метрики
        mandatory_metrics = ['ttd_actions_p90', 'ttd_opp_actions_p90']
        selectable_metrics = [m for m in all_metrics if m not in mandatory_metrics]
        metric_names = {m: METRIC_NAMES_RU.get(m, m) for m in selectable_metrics}
        
        with st.expander("Настройка колонок общей таблицы"):
            # Восстанавливаем сохранённые (без обязательных)
            saved = st.session_state.selected_main_metrics if st.session_state.selected_main_metrics else []
            default_selected = [m for m in saved if m in selectable_metrics]
            selected_optional = st.multiselect(
                "Выберите дополнительные метрики для отображения",
                options=selectable_metrics,
                format_func=lambda m: metric_names[m],
                default=default_selected,
                key="main_metrics_selector_optional"
            )
            selected_metrics = mandatory_metrics + selected_optional
            if selected_optional != saved:
                st.session_state.selected_main_metrics = selected_optional
                save_selected_metrics(selected_optional, SELECTED_METRICS_FILE)
        if not selected_metrics:
            selected_metrics = mandatory_metrics + all_metrics[:3]

        subtabs = st.tabs(["Общий рейтинг", "FW", "AM", "CM", "FB", "CB"])
        with subtabs[0]:
            df_main = build_main_table(df_active, selected_metrics)
            st.dataframe(df_main, height=600, use_container_width=True, on_select="rerun", selection_mode="single-row", key="main_table")
            if "main_table" in st.session_state and st.session_state.main_table.selection.rows:
                idx = next(iter(st.session_state.main_table.selection.rows))
                if idx < len(df_active):
                    player_row = df_active.iloc[idx]
                    col1, col2, col3 = st.columns([1, 2, 1])
                    pos = get_position_group(player_row['position'])
                    weights = st.session_state.current_settings.get(pos, {})
                    sorted_metrics = sorted(weights.items(), key=lambda x: -abs(x[1]))
                    radar_metrics = [m for m, _ in sorted_metrics if m in df_active.columns][:8]
                    if not radar_metrics:
                        radar_metrics = [c for c in df_active.columns if c.endswith('_p90') or c.endswith('_pct')][:8]
                    avg_series = None
                    if st.session_state.avg_source == 'Текущие данные':
                        avg_series = df_active[radar_metrics].mean()
                    elif st.session_state.avg_source == 'Лига из БД' and st.session_state.avg_league and st.session_state.avg_season:
                        df_league = load_from_db([st.session_state.avg_league], [st.session_state.avg_season])
                        if not df_league.empty:
                            df_league = df_league[df_league['minutes'] >= MIN_MINUTES]
                            if not df_league.empty:
                                avg_series = df_league[radar_metrics].mean()
                    with col2:
                        fig = create_player_radar_figure(player_row, df_active, st.session_state.current_settings, avg_values=avg_series)
                        st.plotly_chart(fig, use_container_width=True)
        for i, pos in enumerate(['FW','AM','CM','FB','CB'], 1):
            with subtabs[i]:
                rows, headers = position_tables_active.get(pos, ([], []))
                if rows:
                    numbered_rows = [[j+1] + row for j, row in enumerate(rows)]
                    df_pos = pd.DataFrame(numbered_rows, columns=['№'] + headers)
                    st.dataframe(df_pos, height=400, use_container_width=True, on_select="rerun", selection_mode="single-row", key=f"table_{pos}")
                    state_key = f"table_{pos}"
                    if state_key in st.session_state and st.session_state[state_key].selection.rows:
                        idx = next(iter(st.session_state[state_key].selection.rows))
                        if idx < len(rows):
                            player_name = rows[idx][0]
                            player_min = int(rows[idx][1])
                            candidate = df_active[(df_active['player'] == player_name) & (df_active['minutes'] == player_min)]
                            if not candidate.empty:
                                player_row = candidate.iloc[0]
                                pos_group = get_position_group(player_row['position'])
                                weights = st.session_state.current_settings.get(pos_group, {})
                                sorted_metrics = sorted(weights.items(), key=lambda x: -abs(x[1]))
                                radar_metrics = [m for m, _ in sorted_metrics if m in df_active.columns][:8]
                                if not radar_metrics:
                                    radar_metrics = [c for c in df_active.columns if c.endswith('_p90') or c.endswith('_pct')][:8]
                                avg_series = None
                                if st.session_state.avg_source == 'Текущие данные':
                                    avg_series = df_active[radar_metrics].mean()
                                elif st.session_state.avg_source == 'Лига из БД' and st.session_state.avg_league and st.session_state.avg_season:
                                    df_league = load_from_db([st.session_state.avg_league], [st.session_state.avg_season])
                                    if not df_league.empty:
                                        df_league = df_league[df_league['minutes'] >= MIN_MINUTES]
                                        if not df_league.empty:
                                            avg_series = df_league[radar_metrics].mean()
                                col1, col2, col3 = st.columns([1, 2, 1])
                                with col2:
                                    fig = create_player_radar_figure(player_row, df_active, st.session_state.current_settings, avg_values=avg_series)
                                    st.plotly_chart(fig, use_container_width=True)
                else:
                    st.info(f"Нет игроков позиции {pos}")
        if st.button("📥 Экспорт в Excel", key="export_season"):
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df_main = build_main_table(df_active, selected_metrics)
                df_main.to_excel(writer, sheet_name='Общий рейтинг', index=False)
                for pos, (rows, headers) in position_tables_active.items():
                    if rows:
                        df_pos = pd.DataFrame(rows, columns=headers)
                        df_pos.to_excel(writer, sheet_name=pos, index=False)
                        from openpyxl.formatting.rule import ColorScaleRule
                        from openpyxl.utils import get_column_letter
                        ws = writer.sheets[pos]
                        metric_columns = list(range(4, 4 + len(headers) - 3))
                        for col_idx in metric_columns:
                            col_letter = get_column_letter(col_idx)
                            max_row = len(rows) + 1
                            ws.conditional_formatting.add(
                                f'{col_letter}2:{col_letter}{max_row}',
                                ColorScaleRule(start_type='min', start_color='FFC7CE',
                                              mid_type='percentile', mid_value=50, mid_color='FFFFEB',
                                              end_type='max', end_color='C6EFCE')
                            )
            st.download_button(label="Скачать Excel", data=output.getvalue(), file_name="season_players_rating.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("Загрузите сезонные данные (в боковой панели)")

with tab_match:
    st.markdown("""
    <style>
    .match-table table { font-size: 12px; }
    .match-table th, .match-table td { font-size: 12px; white-space: nowrap; }
    </style>
    """, unsafe_allow_html=True)
    if st.session_state.df_matches:
        match_ids = list(st.session_state.df_matches.keys())
        match_labels = [st.session_state.df_matches[mid]['label'] for mid in match_ids]
        selected_label = st.selectbox("Активный матч", match_labels, key="active_match_selector")
        active_match_id = None
        for mid, data in st.session_state.df_matches.items():
            if data['label'] == selected_label:
                active_match_id = mid
                break
        if active_match_id is not None:
            df_active = st.session_state.df_matches[active_match_id]['df']
            if 'match_tables' in st.session_state and active_match_id in st.session_state.match_tables:
                position_tables_active = st.session_state.match_tables[active_match_id]
            else:
                position_tables_active = build_match_position_tables(df_active, st.session_state.match_settings)
            all_metrics = [m for m in MATCH_ALL_METRICS if m in df_active.columns]
            # Добавляем ТТД-метрики в список доступных, если они есть
            if 'ttd_actions' in df_active.columns:
                all_metrics.append('ttd_actions')
            if 'ttd_opp_actions' in df_active.columns:
                all_metrics.append('ttd_opp_actions')
            mandatory_metrics = ['ttd_actions', 'ttd_opp_actions']
            selectable_metrics = [m for m in all_metrics if m not in mandatory_metrics]
            metric_names = {m: MATCH_METRIC_NAMES_RU.get(m, m) for m in selectable_metrics}
            
            with st.expander("Настройка колонок общей таблицы"):
                saved = st.session_state.match_selected_metrics if st.session_state.match_selected_metrics else []
                default_selected = [m for m in saved if m in selectable_metrics]
                selected_optional = st.multiselect(
                    "Выберите дополнительные метрики для отображения",
                    options=selectable_metrics,
                    format_func=lambda m: metric_names[m],
                    default=default_selected,
                    key="match_main_metrics_selector_optional"
                )
                selected_metrics = mandatory_metrics + selected_optional
                if selected_optional != saved:
                    st.session_state.match_selected_metrics = selected_optional
                    save_match_selected_metrics(selected_optional)
            if not selected_metrics:
                selected_metrics = mandatory_metrics + all_metrics[:3]
            subtabs = st.tabs(["Общий рейтинг", "FW", "AM", "CM", "FB", "CB"])
            st.markdown("""
            <div style="font-size:13px; margin-bottom:8px; line-height:1.6;">
            <b>Обозначения:</b><br>
            🟢 – лучший в команде &nbsp;&nbsp; 🔴 – худший в команде<br>
            🔵 – выше среднего по лиге &nbsp;&nbsp; 🟤 – ниже среднего по лиге<br>
            (↑) – в матче лучше своего среднего за сезон &nbsp;&nbsp; (↓) – хуже своего среднего
            </div>
            """, unsafe_allow_html=True)
            with subtabs[0]:
                league_avg = st.session_state.get('league_avg')
                player_season_map = st.session_state.get('player_season_map')
                df_main = build_match_main_table(df_active, selected_metrics,
                                                league_avg=league_avg,
                                                player_season_map=player_season_map)
                st.write(f'<div class="match-table">{df_main.to_html(escape=False, index=False)}</div>', unsafe_allow_html=True)
            for i, pos in enumerate(['FW','AM','CM','FB','CB'], 1):
                with subtabs[i]:
                    rows, headers = position_tables_active.get(pos, ([], []))
                    if rows:
                        numbered_rows = [[j+1] + row for j, row in enumerate(rows)]
                        df_pos = pd.DataFrame(numbered_rows, columns=['№'] + headers)
                        st.write(f'<div class="match-table">{df_pos.to_html(escape=False, index=False)}</div>', unsafe_allow_html=True)
                    else:
                        st.info(f"Нет игроков позиции {pos}")
            # Экспорт в Excel
            st.markdown("---")
            st.subheader("📥 Экспорт матчей (формат RuStat)")
            col_sel, col_team, col_season = st.columns([2, 1, 1])
            with col_sel:
                export_match_ids = st.multiselect(
                    "Выберите матчи для экспорта",
                    options=list(st.session_state.df_matches.keys()),
                    format_func=lambda mid: st.session_state.df_matches[mid]['label'],
                    key="rustat_export_matches"
                )
            team_name_default = "Динамо Минск"
            season_default = "2025"
            if export_match_ids:
                sample_id = export_match_ids[0]
                sample_df = st.session_state.df_matches[sample_id]['df']
                if not sample_df.empty and 'team' in sample_df.columns:
                    team_name_default = sample_df['team'].iloc[0]
            with col_team:
                export_team = st.text_input("Название команды", value=team_name_default, key="rustat_team")
            with col_season:
                export_season = st.text_input("Сезон (год)", value=season_default, key="rustat_season")
            if export_match_ids and st.button("📥 Экспорт в Excel (RuStat)", key="export_rustat"):
                with st.spinner("Формирование RuStat-отчёта..."):
                    output = io.BytesIO()
                    try:
                        export_matches_advanced(
                            st.session_state.df_matches,
                            export_match_ids,
                            export_team.strip(),
                            export_season.strip(),
                            st.session_state.get('league_avg'),
                            st.session_state.get('player_season_map'),
                            output
                        )
                        output.seek(0)
                        st.download_button(
                            label="Скачать Excel",
                            data=output,
                            file_name=f"RuStat_{export_team}_{export_season}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            key="download_rustat"
                        )
                    except Exception as e:
                        st.error(f"Ошибка: {e}")
            if st.button("📥 Экспорт в Excel (стандартный)", key="export_match"):
                output = io.BytesIO()
                with pd.ExcelWriter(output, engine='openpyxl') as writer:
                    df_main = build_match_main_table(df_active, selected_metrics)
                    df_main.to_excel(writer, sheet_name='Общий рейтинг', index=False)
                    for pos, (rows, headers) in position_tables_active.items():
                        if rows:
                            df_pos = pd.DataFrame(rows, columns=headers)
                            df_pos.to_excel(writer, sheet_name=pos, index=False)
                            from openpyxl.formatting.rule import ColorScaleRule
                            from openpyxl.utils import get_column_letter
                            ws = writer.sheets[pos]
                            metric_columns = list(range(4, 4 + len(headers) - 3))
                            for col_idx in metric_columns:
                                col_letter = get_column_letter(col_idx)
                                max_row = len(rows) + 1
                                ws.conditional_formatting.add(
                                    f'{col_letter}2:{col_letter}{max_row}',
                                    ColorScaleRule(start_type='min', start_color='FFC7CE',
                                                  mid_type='percentile', mid_value=50, mid_color='FFFFEB',
                                                  end_type='max', end_color='C6EFCE')
                                )
                st.download_button(label="Скачать Excel", data=output.getvalue(), file_name="match_players_rating.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("Загрузите матчи (в боковой панели)")

# Остальной код (редакторы весов) – как в предыдущей версии
# ... (здесь должны быть редакторы весов, они не изменились, можно скопировать из предыдущего рабочего кода)
